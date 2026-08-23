from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Multi_Dataset_Diabetes_Model_and_Routing_Report.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5F6B76"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
CAUTION = "7A5A00"
WHITE = "FFFFFF"

USABLE_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS.items():
        tag = "w:start" if side == "start" else "w:end" if side == "end" else f"w:{side}"
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D7DBE2", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    assert sum(widths_dxa) == USABLE_DXA
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run_font(run, *, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None,
                 name: str = "Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph) -> None:
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fallback, fld_end])


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Title": (23, INK, 0, 4),
        "Subtitle": (12, MUTED, 0, 14),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in specs.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = -Inches(0.188)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_section(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run("DIABETES PREDICTION PIPELINE\tTECHNICAL REPORT")
    set_run_font(r, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_field(fp)
    for run in fp.runs:
        set_run_font(run, size=8.5, color=MUTED)


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("TECHNICAL SYSTEM REPORT")
    set_run_font(r, size=9, bold=True, color=CAUTION)

    title = doc.add_paragraph(style="Title")
    title.add_run("Multi-Dataset Diabetes Prediction\nand Model Routing System")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Detailed methodology, evaluation results, mixed-input handling, prediction flow, outputs, and safety controls")

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    r = meta.add_run("Build reviewed: 22 August 2026   |   Intended use: research validation only")
    set_run_font(r, size=9.5, color=MUTED, italic=True)

    callout = doc.add_paragraph()
    callout.paragraph_format.left_indent = Inches(0.12)
    callout.paragraph_format.right_indent = Inches(0.12)
    callout.paragraph_format.space_before = Pt(0)
    callout.paragraph_format.space_after = Pt(12)
    set_paragraph_shading(callout, CALLOUT)
    r = callout.add_run("CURRENT VERDICT  ")
    set_run_font(r, size=10, bold=True, color=CAUTION)
    r = callout.add_run(
        "The routing software is implemented and testable, but the evidence package is not paper-final or clinically deployable. "
        "PIDD and Sylhet are development pilots; BRFSS and Hospital130 are smoke holdouts."
    )
    set_run_font(r, size=10, color=INK)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int],
              numeric_columns: set[int] | None = None, font_size: float = 9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if numeric_columns and idx in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(value)
        set_run_font(r, size=font_size, bold=True, color=INK)

    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if numeric_columns and idx in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=font_size, color="202124")
    set_table_geometry(table, widths)
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)
    return table


def add_callout(doc: Document, label: str, text: str, *, label_color: str = CAUTION) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(10)
    set_paragraph_shading(paragraph, CALLOUT)
    run = paragraph.add_run(f"{label}  ")
    set_run_font(run, size=10, bold=True, color=label_color)
    run = paragraph.add_run(text)
    set_run_font(run, size=10, color=INK)


def add_lead_paragraph(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(f"{label} ")
    set_run_font(run, bold=True, color=INK)
    paragraph.add_run(text)


def add_numbered_list(doc: Document, items: list[str]) -> None:
    """Add visibly stable numbered paragraphs that restart at 1."""
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.32)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        paragraph.add_run(f"{index}.  {item}")


def add_bullet_list(doc: Document, items: list[str]) -> None:
    """Add stable bullet paragraphs without relying on template numbering state."""
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.32)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        paragraph.add_run(f"•  {item}")


def add_body(doc: Document) -> None:
    doc.add_heading("Executive summary", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "The current build is an availability-aware, multi-dataset research pipeline. It does not train one universal model over "
        "columns copied from unrelated cohorts. Instead, it maintains cohort-specific experts for PIDD, Sylhet, BRFSS, and "
        "Hospital130; identifies the largest evaluated feature panel supported by the supplied measurements; applies a panel-specific, "
        "development-selected model and calibrator; and returns a traceable probability or an explicit abstention."
    )
    summary_points = [
        "PIDD, Sylhet, and BRFSS can be present together in one namespaced patient request, but one screening expert supplies the final percentage.",
        "Probabilities are never averaged across those cohorts because there is no matched patient dataset on which such fusion has been trained and validated.",
        "Hospital130 is routed separately because it predicts 30-day readmission, not diabetes status.",
        "Logistic Regression is the benchmark baseline; Random Forest, XGBoost, LightGBM, and CatBoost form the panel-routing model bank.",
        "The percentage is calibrated against the selected cohort's binary label. It is not yet a prospective lifetime-risk or onset-age estimate.",
    ]
    add_bullet_list(doc, summary_points)
    add_callout(
        doc,
        "EVIDENCE STATUS",
        "The software behavior is implemented and the current test suite passes, but the retained studies are development pilots or smoke holdouts. "
        "The publication gate reports 16 blockers, so none of the percentages is approved as a final paper estimate or clinical output.",
    )

    doc.add_heading("1. System scope, datasets, and targets", level=1)
    dataset_rows = [
        ["PIDD", "Clinical/metabolic screening", "Positive diabetes label", "5-fold development pilot", "~34.9% per outer evaluation"],
        ["Sylhet", "Symptom questionnaire", "Positive early-stage diabetes label", "5-fold grouped development pilot", "61.5%"],
        ["BRFSS", "Population health survey", "Self-reported diagnosed diabetes", "10,000-row smoke test", "13.3%"],
        ["Hospital130", "Diabetes-related hospital encounters", "Readmission within 30 days", "10,000-row smoke test", "11.4%"],
    ]
    add_table(
        doc,
        ["Dataset", "Information source", "Target", "Current evaluation", "Observed positive rate"],
        dataset_rows,
        [1200, 2000, 2400, 2280, 1480],
        font_size=8.5,
    )
    add_lead_paragraph(
        doc,
        "Endpoint rule.",
        "A displayed percentage always inherits the endpoint and population of the selected expert. A PIDD percentage is not numerically "
        "interchangeable with a BRFSS percentage, even when both concern diabetes, because their populations, measurements, and target definitions differ."
    )
    add_lead_paragraph(
        doc,
        "Hospital isolation.",
        "Hospital130 answers a different question: whether an encounter will be followed by readmission within 30 days. The mixed diabetes-screening "
        "function rejects it rather than allowing a readmission probability to masquerade as diabetes risk."
    )

    doc.add_heading("2. Input contract and feature-availability panels", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "A request represents exactly one patient. Single-dataset mode accepts one JSON feature object. Mixed mode accepts an object keyed by dataset, "
        "then by the original feature names. Any number of fields may be supplied for each dataset; a dataset becomes eligible only when those fields "
        "complete at least one predeclared panel. Missing values do not count as available."
    )
    add_callout(
        doc,
        "WHY NAMESPACING MATTERS",
        "Feature names are not automatically harmonized. For example, PIDD Age is age in years, whereas BRFSS Age is an encoded age group. "
        "Keeping them under pidd and brfss prevents accidental semantic mixing.",
        label_color=BLUE,
    )

    doc.add_heading("2.1 Exact panel inventory", level=2)
    panel_rows = [
        ["PIDD", "full", "8", "Pregnancies; Glucose; BloodPressure; SkinThickness; Insulin; BMI; DiabetesPedigreeFunction; Age"],
        ["PIDD", "no_insulin_or_skinfold", "6", "Pregnancies; Glucose; BloodPressure; BMI; DiabetesPedigreeFunction; Age"],
        ["PIDD", "routine_clinic", "5", "Pregnancies; Glucose; BloodPressure; BMI; Age"],
        ["PIDD", "core_metabolic", "4", "Glucose; BMI; DiabetesPedigreeFunction; Age"],
        ["PIDD", "non_laboratory", "4", "Pregnancies; BMI; DiabetesPedigreeFunction; Age"],
        ["PIDD", "minimal_glucose_age", "2", "Glucose; Age"],
        ["Sylhet", "full", "16", "Age; Gender; Polyuria; Polydipsia; weight loss; weakness; Polyphagia; thrush; visual blurring; Itching; Irritability; delayed healing; paresis; stiffness; Alopecia; Obesity"],
        ["Sylhet", "core_questionnaire", "10", "Age; Gender; Polyuria; Polydipsia; weight loss; weakness; Polyphagia; visual blurring; delayed healing; Obesity"],
        ["Sylhet", "core_symptoms", "6", "Polyuria; Polydipsia; weight loss; weakness; Polyphagia; delayed healing"],
        ["Sylhet", "classic_symptoms", "4", "Polyuria; Polydipsia; weight loss; Polyphagia"],
        ["Sylhet", "minimal_polyuria_polydipsia", "2", "Polyuria; Polydipsia"],
        ["BRFSS", "full", "21", "Blood pressure/cholesterol; BMI; smoking/stroke/heart disease; activity/diet/alcohol; healthcare access; general, mental, and physical health; walking; sex; age; education; income"],
        ["BRFSS", "routine_survey", "12", "HighBP; HighChol; BMI; Smoker; PhysActivity; GenHlth; PhysHlth; DiffWalk; Sex; Age; Education; Income"],
        ["BRFSS", "cardiometabolic_health", "9", "HighBP; HighChol; BMI; Stroke; HeartDiseaseorAttack; GenHlth; PhysHlth; DiffWalk; Age"],
        ["BRFSS", "health_status", "6", "HighBP; BMI; GenHlth; PhysHlth; DiffWalk; Age"],
        ["BRFSS", "demographics_health", "5", "GenHlth; Sex; Age; Education; Income"],
        ["BRFSS", "minimal_general_health_age", "2", "GenHlth; Age"],
        ["Hospital130", "full", "46", "Age; encounter intensity; prior utilization; admission/discharge/source; race/gender/payer/specialty; three diagnoses; glucose/A1C categories; diabetes medications; change; diabetesMed"],
        ["Hospital130", "encounter_utilization", "9", "Age; time in hospital; lab/procedure/medication counts; outpatient, emergency, and inpatient use; diagnosis count"],
        ["Hospital130", "admission_clinical", "9", "Age; admission/discharge/source; three diagnoses; max glucose; A1C result"],
        ["Hospital130", "prior_utilization", "5", "Age; outpatient; emergency; inpatient; diagnosis count"],
        ["Hospital130", "minimal_inpatient_history", "2", "Age; prior inpatient use"],
    ]
    add_table(
        doc,
        ["Dataset", "Panel", "Fields", "Measurements required"],
        panel_rows,
        [1200, 2500, 700, 4960],
        numeric_columns={2},
        font_size=7.9,
    )
    add_lead_paragraph(
        doc,
        "Panel-selection rule.",
        "If no panel is explicitly requested, the router selects the supported panel with the largest number of fields; full wins a same-size tie, "
        "followed by deterministic scenario-name ordering. It does not create a new ad hoc subset at inference time."
    )

    doc.add_heading("3. Preprocessing, models, and training controls", level=1)
    doc.add_heading("3.1 Fold-fitted preprocessing", level=2)
    preprocessing_points = [
        "Numeric variables are median-imputed within the training fold, with missingness indicators retained.",
        "Logistic Regression additionally receives standardized numeric variables.",
        "Categorical variables are imputed with a dedicated __MISSING__ level and one-hot encoded; unknown categories are ignored rather than crashing inference.",
        "Unselected panel fields are deliberately masked to missing before transformation, reproducing the same feature-availability condition used during routing evaluation.",
        "Preprocessors are fitted only on training/development data and stored with each model bundle.",
    ]
    add_bullet_list(doc, preprocessing_points)

    doc.add_heading("3.2 Model bank", level=2)
    model_rows = [
        ["Logistic Regression", "Benchmark baseline", "Regularized linear log-odds model; numeric scaling; class weighting can be tuned", "Interpretable reference; not a routing candidate"],
        ["Random Forest", "Router candidate", "200-800 trees; depth, leaf size, feature subsampling, and class weighting tuned", "Stable nonlinear ensemble; often selected for sparse panels"],
        ["XGBoost", "Router candidate", "Histogram tree boosting; depth, learning rate, sampling, regularization, and optional positive-class weight tuned", "Strong interaction modeling and ranking performance"],
        ["LightGBM", "Router candidate", "Leaf-wise boosting; leaves, depth, child size, sampling, regularization, and class weighting tuned", "Efficient boosted-tree alternative"],
        ["CatBoost", "Router candidate", "Boosted symmetric trees; depth, learning rate, L2, random strength, and balanced-class option tuned", "Competitive probability performance, especially in current PIDD/Sylhet pilots"],
    ]
    add_table(
        doc,
        ["Model", "Role", "Implementation and tuning", "Interpretation in this pipeline"],
        model_rows,
        [1800, 1500, 3560, 2500],
        font_size=8.2,
    )
    p = doc.add_paragraph()
    p.add_run(
        "Optuna searches the prespecified parameter ranges. The number of trials is intentionally reduced in development-pilot and smoke profiles; "
        "therefore the retained models demonstrate the workflow but do not represent a fully budgeted final comparison. Random seeds and worker counts "
        "are recorded to support repeatability."
    )

    doc.add_heading("3.3 Nested selection, calibration, and thresholds", level=2)
    method_steps = [
        "Create an outer evaluation fold (PIDD/Sylhet) or preserve a locked test partition (BRFSS/Hospital130). Patient-grouped splitting is used when repeated patients exist.",
        "Tune each model only inside development data, using inner cross-validation and the configured optimization metric.",
        "Generate cross-fitted development probabilities for calibration, threshold analysis, and model selection.",
        "Choose calibration only from development predictions. Candidate methods include identity, sigmoid, and—where sample size permits—isotonic calibration.",
        "For each availability panel, select the tree model with the lowest inner cross-fitted Brier score; use log loss and then negative average precision as deterministic tie-breakers.",
        "Evaluate the frozen choice on the corresponding outer fold or locked smoke test. Test labels are not used to select the panel model, calibrator, or threshold.",
    ]
    add_numbered_list(doc, method_steps)
    add_callout(
        doc,
        "THRESHOLD INTERPRETATION",
        "The report retains a continuous calibrated probability. A threshold converts it into a research screening/readmission flag only. "
        "A result below threshold does not rule out the outcome, and a result above threshold is not a diagnosis.",
        label_color=BLUE,
    )

    doc.add_heading("4. Routing architecture", level=1)
    doc.add_heading("4.1 Within-dataset routing", level=2)
    within_steps = [
        "Accept exactly one patient row and determine which configured fields are non-missing.",
        "Identify every predeclared panel whose required fields are fully present.",
        "Select the largest supported panel, unless the caller explicitly requests a named panel.",
        "Load that panel's selected model, fold-fitted preprocessor, and scenario calibrator.",
        "Reindex the row to the dataset contract and mask every field outside the selected panel.",
        "Calculate raw probability, transform it with the saved calibrator, and compare it with the stored threshold.",
        "Across development folds, aggregate calibrated probabilities by their mean and report fold minimum, maximum, and standard deviation.",
    ]
    add_numbered_list(doc, within_steps)

    doc.add_heading("4.2 Cross-dataset routing for mixed input", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "PIDD, Sylhet, and BRFSS are evaluated independently. A candidate summary is created for every dataset that receives fields. If only one expert "
        "can predict, it is selected unambiguously. If multiple experts can predict, the default is abstention because there is no matched multi-environment "
        "validation cohort. The caller can supply --prefer-dataset when the appropriate measurement environment is known."
    )
    selection_rows = [
        ["Only one eligible expert", "Select it", "Unambiguous within current software contract"],
        ["Several eligible; --prefer-dataset supplied", "Select requested eligible expert", "Explicit environment choice; still inherits that study's evidence tier"],
        ["Several eligible; no preference", "Abstain", "Automatic cross-dataset selector is not validated"],
        ["Several eligible; experimental flag", "Use PIDD > Sylhet > BRFSS priority, then panel coverage", "Software demonstration only; excluded from efficacy claims"],
        ["Hospital130 included in screening request", "Reject", "Endpoint mismatch: readmission is not diabetes screening"],
    ]
    add_table(
        doc,
        ["Situation", "Behavior", "Rationale"],
        selection_rows,
        [2600, 2500, 4260],
        font_size=8.5,
    )
    add_callout(
        doc,
        "NON-FUSION GUARANTEE",
        "Fields from unselected datasets are listed as not used in the selected percentage. Candidate probabilities may be shown for auditability, "
        "but they are never averaged, multiplied, or otherwise combined.",
    )

    doc.add_heading("5. Model results", level=1)
    doc.add_heading("5.1 Full-input calibrated model comparison", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "AUROC measures ranking discrimination. Average precision (AP) emphasizes performance on positive cases and is prevalence-sensitive. Brier score "
        "is mean squared probability error, while log loss penalizes confident wrong probabilities. Lower Brier and log loss are better. Values below are "
        "calibrated fixed-0.5 summaries; model selection itself is performed on development evidence, not on these displayed evaluation outcomes."
    )
    result_rows = [
        ["PIDD", "CatBoost", "0.8404", "0.7205", "0.1541", "0.4688"],
        ["PIDD", "XGBoost", "0.8389", "0.7275", "0.1547", "0.4693"],
        ["PIDD", "Logistic Regression", "0.8375", "0.7258", "0.1563", "0.4775"],
        ["PIDD", "Random Forest", "0.8370", "0.7086", "0.1558", "0.4736"],
        ["PIDD", "LightGBM", "0.8352", "0.7168", "0.1553", "0.4719"],
        ["Sylhet", "CatBoost", "0.9709", "0.9848", "0.0642", "0.2093"],
        ["Sylhet", "LightGBM", "0.9642", "0.9808", "0.0701", "0.2373"],
        ["Sylhet", "XGBoost", "0.9621", "0.9787", "0.0800", "0.2665"],
        ["Sylhet", "Random Forest", "0.9614", "0.9794", "0.0759", "0.2518"],
        ["Sylhet", "Logistic Regression", "0.9576", "0.9777", "0.0871", "0.3026"],
        ["BRFSS", "XGBoost", "0.8318", "0.4221", "0.0935", "0.3044"],
        ["BRFSS", "CatBoost", "0.8308", "0.4199", "0.0936", "0.3039"],
        ["BRFSS", "Random Forest", "0.8257", "0.4148", "0.0940", "0.3073"],
        ["BRFSS", "LightGBM", "0.8135", "0.3936", "0.0963", "0.3142"],
        ["Hospital130", "XGBoost", "0.6589", "0.2065", "0.0973", "0.3382"],
        ["Hospital130", "Random Forest", "0.6584", "0.2040", "0.0973", "0.3383"],
        ["Hospital130", "CatBoost", "0.6516", "0.2003", "0.0976", "0.3396"],
        ["Hospital130", "LightGBM", "0.6334", "0.1874", "0.0984", "0.3430"],
    ]
    add_table(
        doc,
        ["Dataset", "Model", "AUROC", "AP", "Brier", "Log loss"],
        result_rows,
        [1400, 2240, 1430, 1430, 1430, 1430],
        numeric_columns={2, 3, 4, 5},
        font_size=8.7,
    )

    doc.add_heading("5.2 Availability-router performance by panel", level=2)
    panel_metric_rows = [
        ["PIDD", "minimal_glucose_age", "2", "CB 2; LGBM 1; RF 1; XGB 1", "0.8107", "0.6941", "0.1676", "0.5055"],
        ["PIDD", "core_metabolic", "4", "XGB 2; CB/LGBM/RF 1 each", "0.8360", "0.7279", "0.1551", "0.4706"],
        ["PIDD", "non_laboratory", "4", "XGB 3; LGBM 1; RF 1", "0.7484", "0.6005", "0.1913", "0.5626"],
        ["PIDD", "routine_clinic", "5", "CB 3; RF 1; XGB 1", "0.8378", "0.7155", "0.1563", "0.4756"],
        ["PIDD", "no_insulin_or_skinfold", "6", "CB 3; RF 2", "0.8401", "0.7290", "0.1540", "0.4683"],
        ["PIDD", "full", "8", "CB 3; RF 2", "0.8407", "0.7205", "0.1540", "0.4686"],
        ["Sylhet", "minimal_polyuria_polydipsia", "2", "RF 2; XGB 2; LGBM 1", "0.9052", "0.9314", "0.1126", "0.3620"],
        ["Sylhet", "classic_symptoms", "4", "RF 4; XGB 1", "0.9193", "0.9514", "0.1071", "0.3459"],
        ["Sylhet", "core_symptoms", "6", "RF 2; XGB 2; LGBM 1", "0.9148", "0.9548", "0.1083", "0.3501"],
        ["Sylhet", "core_questionnaire", "10", "XGB 3; RF 2", "0.9471", "0.9719", "0.0908", "0.3055"],
        ["Sylhet", "full", "16", "CB 4; XGB 1", "0.9650", "0.9815", "0.0741", "0.2325"],
        ["BRFSS", "minimal_general_health_age", "2", "RF", "0.7770", "0.3018", "0.1023", "0.3358"],
        ["BRFSS", "demographics_health", "5", "RF", "0.7788", "0.3164", "0.1021", "0.3349"],
        ["BRFSS", "health_status", "6", "XGB", "0.8206", "0.3922", "0.0961", "0.3123"],
        ["BRFSS", "cardiometabolic_health", "9", "XGB", "0.8280", "0.4179", "0.0942", "0.3070"],
        ["BRFSS", "routine_survey", "12", "XGB", "0.8276", "0.4178", "0.0943", "0.3072"],
        ["BRFSS", "full", "21", "CB", "0.8308", "0.4199", "0.0936", "0.3039"],
        ["Hospital130", "minimal_inpatient_history", "2", "RF", "0.6191", "0.1823", "0.0990", "0.3467"],
        ["Hospital130", "prior_utilization", "5", "RF", "0.6114", "0.1815", "0.0992", "0.3472"],
        ["Hospital130", "admission_clinical", "9", "RF", "0.5820", "0.1665", "0.0999", "0.3499"],
        ["Hospital130", "encounter_utilization", "9", "RF", "0.6203", "0.1830", "0.0990", "0.3462"],
        ["Hospital130", "full", "46", "RF", "0.6584", "0.2040", "0.0973", "0.3383"],
    ]
    add_table(
        doc,
        ["Dataset", "Panel", "N", "Selected model(s)", "AUROC", "AP", "Brier", "Log loss"],
        panel_metric_rows,
        [1050, 2060, 500, 2150, 900, 900, 900, 900],
        numeric_columns={2, 4, 5, 6, 7},
        font_size=7.6,
    )
    p = doc.add_paragraph()
    r = p.add_run("Abbreviations: ")
    set_run_font(r, bold=True, color=INK)
    p.add_run("CB = CatBoost; RF = Random Forest; XGB = XGBoost; LGBM = LightGBM. PIDD/Sylhet counts describe selections across five outer-fold routers; BRFSS/Hospital130 have one development-locked smoke router.")

    doc.add_heading("5.3 Result interpretation", level=2)
    interpretations = [
        "PIDD: the six-field panel without insulin or skinfold retains almost all full-panel performance. The non-laboratory panel loses the most discrimination, showing that BMI and demographic history do not substitute for glucose measurement.",
        "Sylhet: more questionnaire fields generally improve performance, but the 520-row symptom-driven cohort is small. The full-panel CatBoost calibration intercept and slope are extreme, indicating unstable probability calibration even though ranking metrics are high.",
        "BRFSS: health-status and cardiometabolic fields provide most of the smoke-run discrimination. A fixed 0.5 threshold has very low sensitivity in this low-prevalence cohort, so probability and development-selected threshold analysis are more informative than accuracy alone.",
        "Hospital130: all panels have modest discrimination, and the full encounter record performs best. This task should be reported independently from diabetes-screening results.",
        "Dynamic switching: selection patterns vary by fold, but gains over strong fixed-model baselines are small or unstable. The protocol therefore treats model switching as a secondary negative result rather than the paper's primary novelty claim.",
    ]
    add_bullet_list(doc, interpretations)

    doc.add_heading("6. Worked mixed-patient example", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "The retained synthetic example supplies two fields from each of three screening datasets. It demonstrates software routing only and is not a clinical case."
    )
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.right_indent = Inches(0.2)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(10)
    set_paragraph_shading(code, LIGHT_GRAY)
    run = code.add_run(
        '{\n  "pidd": {"Glucose": 142, "Age": 34},\n'
        '  "sylhet": {"Polyuria": 1, "Polydipsia": 0},\n'
        '  "brfss": {"GenHlth": 3, "Age": 5}\n}'
    )
    set_run_font(run, size=9, name="Courier New", color="202124")
    candidate_rows = [
        ["PIDD", "minimal_glucose_age", "Glucose, Age", "55.2%", "Eligible"],
        ["Sylhet", "minimal_polyuria_polydipsia", "Polyuria, Polydipsia", "82.9%", "Eligible"],
        ["BRFSS", "minimal_general_health_age", "GenHlth, coded Age", "8.7%", "Eligible"],
    ]
    add_table(
        doc,
        ["Candidate expert", "Panel", "Fields used", "Candidate percentage", "Status"],
        candidate_rows,
        [1250, 2700, 2400, 1700, 1310],
        numeric_columns={3},
        font_size=8.5,
    )
    p = doc.add_paragraph()
    p.add_run(
        "Because all three experts are eligible, the default mixed call would abstain. In the saved example, PIDD was explicitly preferred, so the displayed "
        "percentage is 55.2% from the PIDD glucose-and-age panel. Sylhet and BRFSS fields are listed as not used in that selected percentage. The three "
        "candidate values are not votes and are not contradictory estimates of one universally defined risk; they belong to different cohort models."
    )
    add_callout(
        doc,
        "HOW TO READ 55.2%",
        "It means the calibrated PIDD research ensemble assigned an average probability of 0.552 to the PIDD positive label for this synthetic input. "
        "Four of five fold models were at or above the stored 0.5 research threshold. The fold range was approximately 47.1% to 60.8%, with SD 5.0 percentage points. "
        "That spread is model-fold variation, not a patient confidence interval.",
        label_color=BLUE,
    )

    doc.add_heading("7. Output schema and interpretation", level=1)
    output_rows = [
        ["status / reason", "predicted or abstain; machine-readable abstention code when unavailable", "Whether a percentage was validly produced"],
        ["routing", "selected dataset, endpoint, selection reason, candidate summaries, priority metadata", "Which cohort expert supplied the displayed result and why"],
        ["probability_result", "mean calibrated probability, percentage, fold min/max/SD, threshold, folds above threshold", "Continuous estimate plus model-fold behavior"],
        ["feature_availability", "selected panel, available fields, fields used, panel-excluded fields, unknown fields", "Exactly which measurements entered the model"],
        ["model_provenance", "model and calibrator counts plus per-fold raw/calibrated probabilities", "Traceability of model selection and calibration"],
        ["explanation", "top SHAP factors, directions, normalized importance, additivity checks", "How observed fields changed each fold model's output"],
        ["cohort_context", "unadjusted prevalence in reference-data bins for eligible numeric factors", "Descriptive context only; not causal or a cutoff"],
        ["applicability", "reference-range status, warnings, and cohort-specific notes", "Whether numeric inputs remain within observed support"],
        ["research_evidence", "study profile, evidence tier, selector validation, publication eligibility", "Whether the result may enter a paper claim"],
        ["prediction_targets", "diabetes_probability plus reserved onset_age target", "Stable interface for current and future targets"],
        ["pipeline_source_tree_sha256", "hash of the source tree", "Links the output to the implementation version"],
    ]
    add_table(
        doc,
        ["Output section", "What it contains", "Correct interpretation"],
        output_rows,
        [2400, 3710, 3250],
        font_size=7.9,
    )
    doc.add_heading("7.1 Explanation safeguards", level=2)
    explanation_points = [
        "SHAP values are computed on the underlying fold model and checked for additivity; a failed reconstruction raises an error.",
        "Contributions are grouped back to original source features after one-hot encoding and missingness indicators.",
        "Because model families use different raw-output scales, absolute SHAP contributions are normalized within each fold before aggregation.",
        "The report labels whether a factor mostly raises, mostly lowers, or behaves inconsistently across folds.",
        "SHAP is an explanation of model behavior, not evidence that changing the feature would cause a change in disease risk.",
    ]
    add_bullet_list(doc, explanation_points)

    doc.add_heading("8. Abstention and failure behavior", level=1)
    abstention_rows = [
        ["invalid_input", "A configured numeric value cannot be converted or is non-finite", "Return offending field and provided value"],
        ["no_supported_feature_panel", "Available fields do not complete any configured panel", "Return the available fields; no probability"],
        ["required_scenario_features_missing", "Caller requests a scenario but omits one or more required fields", "Return scenario and missing fields"],
        ["outside_reference_range", "A used numeric input falls below/above the reference cohort range", "Abstain by default; --allow-out-of-range changes this to a warning"],
        ["no_supported_dataset_panel", "No dataset-specific object produces an eligible result", "Return each dataset attempt and its reason"],
        ["cross_dataset_environment_selection_not_validated", "More than one screening expert is eligible without an explicit preference", "Return candidate summaries and require appropriate environment selection"],
        ["endpoint mismatch", "Hospital130 or another non-screening task is supplied to the mixed screening function", "Reject the request rather than relabel the target"],
        ["bundle/provenance error", "Bundles disagree on feature/scenario contracts or required objects are missing", "Raise an error; never silently fall back"],
    ]
    add_table(
        doc,
        ["Reason", "Trigger", "System behavior"],
        abstention_rows,
        [3150, 3200, 3010],
        font_size=7.7,
    )
    p = doc.add_paragraph()
    p.add_run(
        "Unknown extra input fields are recorded and ignored for a single-dataset prediction. Unknown dataset namespaces, an empty mixed request, a preferred "
        "dataset that did not produce an eligible prediction, or more than one patient row are treated as input errors. Patient data are not persisted unless "
        "an output path is explicitly supplied."
    )

    doc.add_heading("9. Research validation and publication readiness", level=1)
    doc.add_heading("9.1 Software verification", level=2)
    verification_rows = [
        ["Automated tests", "36 tests passed; 5 subtests passed", "Covers routing, patient results, readiness controls, and supporting pipeline behavior"],
        ["Bundle provenance", "Experiment IDs, configuration hashes, raw-data hashes, and source-tree hash retained", "Supports traceability and mismatch detection"],
        ["Leakage controls", "Fold-fitted preprocessing; inner development selection; locked smoke-test evaluation", "Prevents evaluation labels from selecting models/calibrators"],
        ["Explanation validation", "Per-fold SHAP additivity check", "Detects incompatible explanation output"],
        ["Publication gate", "Automated verdict: NOT READY; 16 blockers and 1 warning", "Prevents pilot/smoke outputs from being mislabeled as final evidence"],
    ]
    add_table(doc, ["Control", "Current result", "Purpose"], verification_rows, [1900, 3100, 4360], font_size=8.4)

    doc.add_heading("9.2 What remains before a defensible paper", level=2)
    remaining_rows = [
        ["Final experiments", "Run prespecified final profiles; use full BRFSS; freeze all choices before test access"],
        ["Protocol and claims", "Freeze protocol; complete novelty review and claim-to-evidence table"],
        ["Probability evaluation", "Complete calibration plots/intervals, decision curves, and clinical-utility interpretation"],
        ["Robustness", "Complete ablations, risk-coverage/selective-prediction analysis, and subgroup/fairness evaluation"],
        ["Reporting", "Complete TRIPOD+AI checklist and PROBAST+AI risk-of-bias/applicability assessment"],
        ["Transportability", "Add external or temporal validation before journal-strength or clinical-generalization claims"],
        ["Mixed selector", "Obtain matched multi-environment patient data before validating automatic cross-dataset selection or fusion"],
    ]
    add_table(doc, ["Workstream", "Required completion"], remaining_rows, [2200, 7160], font_size=8.5)
    add_callout(
        doc,
        "DEFENSIBLE PRESENT CLAIM",
        "The system can quantify how discrimination and calibration change across prespecified feature-availability environments and can return traceable, "
        "cohort-specific probabilities with explicit abstention. The current evidence does not support clinical deployment or a headline claim that dynamic "
        "algorithm switching is superior.",
    )

    doc.add_heading("10. Future onset-age extension", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "The schema separates diabetes_probability from onset_age. onset_age currently returns status = not_implemented and a null age, preventing the "
        "classifier percentage from being misrepresented as a prediction of when disease will occur."
    )
    p = doc.add_paragraph()
    run = p.add_run("Data and model requirement. ")
    set_run_font(run, bold=True, color=INK)
    p.add_run(
        "A validated target requires longitudinal follow-up or reliable onset-time labels, which the present cross-sectional screening cohorts do not provide. "
        "A future survival/time-to-event model must handle censoring and be validated separately for its inputs, population, calibration horizon, uncertainty, and abstention. "
    )
    run = p.add_run("Compatibility. ")
    set_run_font(run, bold=True, color=INK)
    p.add_run("It can populate prediction_targets.onset_age without changing diabetes_probability or the current API contract. ")
    run = p.add_run("Conclusion. ")
    set_run_font(run, bold=True, color=BLUE)
    p.add_run(
        "The four-pipeline router is a traceable research prototype; publication claims still require locked experiments, external validation, and bias/reporting work."
    )

def set_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Multi-Dataset Diabetes Prediction and Model Routing System"
    props.subject = "Models, results, input/output behavior, routing, and abstention rules"
    props.author = "Diabetes Prediction Research Project"
    props.keywords = "diabetes, routing, PIDD, Sylhet, BRFSS, Hospital130, calibration"
    props.comments = "Research-validation report generated from repository evidence artifacts."


def main() -> None:
    doc = Document()
    configure_styles(doc)
    configure_section(doc)
    set_document_properties(doc)
    add_title_block(doc)
    add_body(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
